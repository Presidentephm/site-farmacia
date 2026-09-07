# -*- coding: utf-8 -*-
"""Dias trabalhados e vendas da ANA CELIA (folguista) na loja Centro - agosto/2026.
Fonte: InovaFarma, relatorio detalhado por produto, extraido em 03/09/2026."""
import openpyxl, re, datetime
from collections import defaultdict
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SRC = "/root/.claude/uploads/f4fc9be7-f091-55cf-869c-2f6961829939/6e9cc1cb-Comiss_o_de_Vendedores0309centro.xlsx"
OUT = "/home/user/site-farmacia/relatorios/contabilidade/ANA_CELIA_DIAS_TRABALHADOS_AGOSTO_2026.xlsx"
VENDEDOR, COD, DIARIA = "ANA CELIA", "131", 100.00

F = "Arial"; MONEY = 'R$ #,##0.00;-R$ #,##0.00;"-"'
def font(sz=10, b=False, color="000000", it=False):
    return Font(name=F, size=sz, bold=b, color=color, italic=it)
FILL_TIT = PatternFill("solid", fgColor="1F3864")
FILL_HDR = PatternFill("solid", fgColor="2E5496")
FILL_SEC = PatternFill("solid", fgColor="D9E2F3")
FILL_TOT = PatternFill("solid", fgColor="E2EFDA")
FILL_LIQ = PatternFill("solid", fgColor="C6E0B4")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_DOM = PatternFill("solid", fgColor="FFF0F0")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BLUE = "0000FF"

wb_src = openpyxl.load_workbook(SRC, data_only=True)
ws_src = wb_src[wb_src.sheetnames[0]]
rows = [[('' if c is None else str(c).strip()) for c in r] for r in ws_src.iter_rows(values_only=True)]
cur = None
dias = defaultdict(lambda: dict(ped=set(), itens=0.0, bruta=0.0, liq=0.0, com=0.0, inc=0.0,
                                hmin="99:99", hmax="00:00"))
for r in rows:
    v = [c for c in r if c]
    if len(v) == 2 and re.fullmatch(r"\d+", v[0]):
        cur = v[1]; continue
    if len(v) >= 12 and re.fullmatch(r"\d+", v[0]) and cur == VENDEDOR:
        ped, data, _cli, _p, _d, qtd, pb, _pd, liq, _pc, vcom, inc = v[:12]
        d = data[:10]; h = data[11:16]
        x = dias[d]
        x["ped"].add(ped); x["itens"] += float(qtd)
        x["bruta"] += float(pb) * float(qtd); x["liq"] += float(liq)
        x["com"] += float(vcom); x["inc"] += float(inc)
        x["hmin"] = min(x["hmin"], h); x["hmax"] = max(x["hmax"], h)

SEM = ["segunda", "terça", "quarta", "quinta", "sexta", "sábado", "domingo"]
wb = openpyxl.Workbook()
ws = wb.active; ws.title = "ANA CELIA AGO.26"
ws.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHI", [13, 12, 9, 9, 14, 14, 12, 11, 16]):
    ws.column_dimensions[col].width = w
ws.merge_cells("A1:I1")
c = ws.cell(1, 1, "DIAS TRABALHADOS E VENDAS — ANA CELIA (FOLGUISTA)")
c.font = font(13, True, "FFFFFF"); c.fill = FILL_TIT
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 26
ws.merge_cells("A2:I2")
c = ws.cell(2, 1, f"Loja CENTRO (matriz) · código {COD} · agosto/2026 · fonte InovaFarma, extraído em 03/09/2026")
c.font = font(10, True, "1F3864"); c.alignment = Alignment(horizontal="center")

hdr = ["DATA", "DIA DA SEMANA", "PEDIDOS", "ITENS", "VENDA BRUTA", "VENDA LÍQUIDA",
       "COMISSÃO", "INCENTIVO", "1ª / ÚLTIMA VENDA"]
for i, h in enumerate(hdr, 1):
    c = ws.cell(4, i, h); c.font = font(9, True, "FFFFFF"); c.fill = FILL_HDR
    c.alignment = Alignment(horizontal="center", wrap_text=True); c.border = BOX
ws.row_dimensions[4].height = 28
ws.freeze_panes = "A5"
r = 5; ini = r
for d in sorted(dias):
    x = dias[d]
    dt = datetime.date(*map(int, d.split("-")))
    ws.cell(r, 1, dt.strftime("%d/%m/%Y")).font = font(10, True)
    ws.cell(r, 2, SEM[dt.weekday()]).font = font(10)
    ws.cell(r, 3, len(x["ped"])).alignment = Alignment(horizontal="center")
    ws.cell(r, 4, int(x["itens"])).alignment = Alignment(horizontal="center")
    for i, v in ((5, x["bruta"]), (6, x["liq"]), (7, x["com"]), (8, x["inc"])):
        cc = ws.cell(r, i, round(v, 2)); cc.number_format = MONEY
    ws.cell(r, 9, f'{x["hmin"]} — {x["hmax"]}').alignment = Alignment(horizontal="center")
    for i in range(1, 10):
        ws.cell(r, i).border = BOX
        if dt.weekday() == 6:
            ws.cell(r, i).fill = FILL_DOM
    r += 1
fim = r - 1
ws.cell(r, 1, "TOTAL").font = font(10, True)
ws.cell(r, 2, f"{len(dias)} dias").font = font(10, True)
for i in (3, 4, 5, 6, 7, 8):
    L = openpyxl.utils.get_column_letter(i)
    cc = ws.cell(r, i, f"=SUM({L}{ini}:{L}{fim})")
    cc.number_format = MONEY if i >= 5 else "0"
    cc.font = font(10, True)
    if i < 5:
        cc.alignment = Alignment(horizontal="center")
for i in range(1, 10):
    ws.cell(r, i).fill = FILL_TOT; ws.cell(r, i).border = BOX
tot = r
r += 2

for i in range(1, 10):
    ws.cell(r, i).fill = FILL_SEC
ws.cell(r, 1, "DIÁRIAS PARA O CONTAS A PAGAR").font = font(11, True, "1F3864")
r += 1
def bloco(rot, val, obs="", fill=None, bold=False, fmt=MONEY, cor=None):
    global r
    c = ws.cell(r, 1, rot); c.font = font(10, bold); c.border = BOX
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    c2 = ws.cell(r, 5, val); c2.number_format = fmt
    c2.font = font(11 if bold else 10, bold, cor or "000000"); c2.border = BOX
    if fill:
        for i in range(1, 6):
            ws.cell(r, i).fill = fill
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=9)
    c3 = ws.cell(r, 6, obs); c3.font = font(9, it=True)
    c3.alignment = Alignment(wrap_text=True, vertical="center")
    r += 1
    return r - 1
l_dias = bloco("Dias com venda registrada", len(dias), "Cada dia com pedido no código dela conta como um dia trabalhado.", fmt="0")
l_diaria = bloco("Valor da diária", DIARIA, "Valor combinado com a empresa.", cor=BLUE, fill=FILL_IN)
bloco("TOTAL DE DIÁRIAS", f"=E{l_dias}*E{l_diaria}", "Levar para o contas a pagar / aba FOLGUISTAS do relatório do Centro.",
      fill=FILL_LIQ, bold=True)
bloco("Comissão apurada no período", f"=G{tot}", "Confirmar se folguista recebe comissão além da diária.")
r += 1

notas = ["COMO LER ESTE RELATÓRIO:",
         "A lista traz os dias em que houve VENDA registrada no código 131. Um dia em que ela tenha trabalhado sem realizar nenhuma venda não aparece aqui — conferir com a escala da loja.",
         "Os horários da última coluna são da primeira e da última venda do dia; servem para conferir o turno, não são registro de ponto.",
         "Os domingos estão destacados em vermelho claro.",
         "Ela trabalhou sempre nas quartas, sextas e domingos — cobrindo as folgas da equipe. A única sexta sem venda no mês foi 21/08.",
         "Estes valores NÃO entram no holerite: folguista é pagamento por diária, lançado no contas a pagar."]
for t in notas:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    c = ws.cell(r, 1, t)
    c.font = font(9, t.endswith(":"), "C00000" if t.endswith(":") else "000000")
    c.alignment = Alignment(wrap_text=True, vertical="center")
    r += 1
ws.auto_filter.ref = f"A4:I{fim}"
wb.save(OUT)
print("ok", OUT, "| dias:", len(dias))
